from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

import os, sys
if __name__ == "__main__":
    sys.path.append(os.getcwd())


class Logger(object):
    logger_dict = {
        "likelihood"  : "likelihood_logger",
        "prediction"  : "prediction_logger_36M",
        "prediction_Eva"  : "prediction_logger_Eva",
        "prediction_AMASS"  : "prediction_logger_AMASS",
        "calibration" : "calibration_logger",
        "MAE" : "MAE"
    }
    
    alignment_type = Alignment(horizontal="left", vertical="center", wrap_text=True)
    font_type = Font(name="Calibri")
    
    @staticmethod
    def open(workbook, sheet):
        try:
            wb = load_workbook(filename=workbook)
            sh = wb[Logger.logger_dict[sheet]]
        except IOError:
            print("ERROR: CAN NOT OPEN %s" % (workbook))
            wb, sh = None, None
        finally:
            return wb, sh
    
    @staticmethod
    def find(sheet, key):
        for cell in iter(sheet[1]):
            if cell.value == key:
                return cell.row, cell.column
        return 1, cell.column + 1            
        
    @staticmethod
    def write(workbook:str, sheet:str, overwrite:bool = False, data:dict = {}):
        wb, sh = Logger.open(workbook, sheet)
        if not wb or not sh:
            print("ERROR: There is no workbook/sheet exists.")
        else:            
            row = len(sh['B']) if overwrite else len(sh['B'])+1
            for key, value in data.items():
                t_row, t_col = Logger.find(sh, key)
                if sh.cell(row=t_row, column=t_col).value is None:
                    sh.cell(row=t_row, column=t_col).alignment = Logger.alignment_type
                    sh.cell(row=t_row, column=t_col).font = Logger.font_type
                    sh.cell(row=t_row, column=t_col).value = key
                
                column = t_col
                sh.cell(row=t_row, column=t_col).alignment = Logger.alignment_type
                sh.cell(row=t_row, column=t_col).font = Logger.font_type
                sh.cell(row=row, column=column).value = value
            wb.save(filename=workbook)



if __name__ == "__main__":
    logger = "logger.xlsx"
    Logger.write(workbook=logger, sheet='prediction', overwrite=False, data={"test": 1})
                